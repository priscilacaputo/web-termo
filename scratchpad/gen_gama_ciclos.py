# -*- coding: utf-8 -*-
"""Genera assets/js/gama-ciclos-data.js: datos para saber la GAMA (mensual, trimestral,
anual...) de cada OT preventiva en Programación de OTs.

  GAMA_HR[ "RUTA/CONT" ]  = ciclos (en días) de los paquetes que usa esa hoja de ruta.
                            Si es uno solo, toda OT de esa hoja de ruta es de esa gama.
  GAMA_POS[ posición ]    = [offsets, base, nRef, fechaRef, "RUTA/CONT"] para posiciones con
                            hoja de ruta de varios paquetes cuyo punto del ciclo se pudo deducir
                            del historial (lista de operaciones IW49 + tomas IP24).
    SAP (estrategia día fijo, sin offsets de paquete) arma la toma n con los paquetes cuyo
    ciclo divide a  base·n + offset  (en meses); el offset depende de dónde arrancó cada plan.
    Se guardan los offsets que el historial deja posibles; la página da la gama solo si todos
    coinciden para esa toma (validado leave-one-out: 0 errores en 442 casos, 2026-09-22).

Uso: python scratchpad/gen_gama_ciclos.py
Fuentes (Descargas): HDR ter aep.xlsm (IA17), IP24 (1).xlsx, lista de operaciones.xlsx,
EXPORT.XLSX (IP11 estrategias, hoja Paquetes).
"""
import os, re, json, collections, datetime as dt
import openpyxl

D = os.path.join(os.path.expanduser('~'), 'Downloads')
OUT = os.path.join(os.path.dirname(__file__), '..', 'assets', 'js', 'gama-ciclos-data.js')

# ── paquetes de cada estrategia: número de paquete → ciclo en días ──
paq = collections.defaultdict(dict)
for r in openpyxl.load_workbook(os.path.join(D, 'EXPORT.XLSX'), read_only=True, data_only=True)['Paquetes'].iter_rows(min_row=2, values_only=True):
    if not r or not r[0]:
        continue
    seg = float(r[4])
    dias = round(seg / 86400)
    unidad = str(r[3])
    if unidad in ('MON', 'JHR'):                       # meses / años "comerciales": 30 / 360 días
        dias = round(seg / 2592000) * 30
    paq[str(r[0]).strip()][str(r[1]).strip().lstrip('0') or '0'] = dias

# ── IP24: tomas por posición y estrategia de cada hoja de ruta ──
tomas = collections.defaultdict(list)
pos_de_orden = {}
estr_hr = {}
hr_pos = {}
for r in openpyxl.load_workbook(os.path.join(D, 'IP24 (1).xlsx'), read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True):
    pos, estr, n, f, orden, ruta, cont = str(r[1]), r[3], r[5], r[6], str(r[7] or ''), str(r[8] or ''), str(r[9] or '')
    hr = f'{ruta}/{cont}'
    tomas[pos].append((n, f, orden))
    hr_pos[pos] = hr
    estr_hr.setdefault(hr, estr)
    if orden:
        pos_de_orden[orden] = (pos, n)

# ── IA17: paquetes de cada hoja de ruta ──
GAMA_HR = collections.defaultdict(set)
for r in openpyxl.load_workbook(os.path.join(D, 'HDR ter aep.xlsm'), data_only=True, read_only=True)['Hojas de Ruta'].iter_rows(min_row=2, values_only=True):
    if not r or not r[0] or r[12] in (None, ''):
        continue
    hr = f'{str(r[0]).strip()}/{str(r[1]).strip()}'
    if hr not in estr_hr:          # hoja de ruta que ningún plan usa: sin estrategia, no se puede traducir
        continue
    estr = estr_hr[hr]
    k = str(r[12]).strip().lstrip('0') or '0'
    if k in paq.get(estr, {}):
        GAMA_HR[hr].add(paq[estr][k])

# ── historial: gama observada de OTs pasadas (texto de las operaciones) ──
MESES = [('Mensual', 1, r'mensual|(^|[^0-9])1\s*m\b'), ('Bimestral', 2, r'bimestral|(^|[^0-9])2\s*m\b'),
         ('Trimestral', 3, r'trimestral|(^|[^0-9])3\s*m\b'), ('Cuatrimestral', 4, r'cuatrimestral|(^|[^0-9])4\s*m\b'),
         ('Semestral', 6, r'semestral|(^|[^0-9])6\s*m\b'), ('Anual', 12, r'anual|(^|[^0-9])1\s*a\b')]
ops = collections.defaultdict(list)
for r in openpyxl.load_workbook(os.path.join(D, 'lista de operaciones.xlsx'), read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True):
    ops[str(r[1])].append(str(r[3] or '').lower())
obs = collections.defaultdict(list)
for o, ts in ops.items():
    ms = [m for _, m, rx in MESES for t in ts if re.search(rx, t)]
    if ms and o in pos_de_orden:
        pos, n = pos_de_orden[o]
        obs[pos].append((n, max(ms)))


def due(m, cyc):
    ds = [c for c in cyc if m % c == 0]
    return max(ds) if ds else None


GAMA_POS = {}
for pos, ob in obs.items():
    hr = hr_pos.get(pos)
    cyc_d = GAMA_HR.get(hr, set())
    if len(cyc_d) < 2 or any(c % 30 for c in cyc_d):
        continue
    cyc = sorted(c // 30 for c in cyc_d)               # meses
    b = cyc[0]
    offs = [off for off in range(0, max(cyc), b) if all(due(b * n + off, cyc) == g for n, g in ob)]
    if not offs or len(offs) == len(range(0, max(cyc), b)):   # historial inconsistente, o no descarta nada
        continue
    ref = max(((n, f) for n, f, _ in tomas[pos] if isinstance(f, dt.datetime)), default=None)
    if not ref:
        continue
    GAMA_POS[pos] = [offs, b, ref[0], ref[1].strftime('%Y-%m-%d'), hr]

hdr = {k: sorted(v) for k, v in sorted(GAMA_HR.items())}
with open(OUT, 'w', encoding='utf-8') as fh:
    fh.write('/* ─── GAMA_HR / GAMA_POS — para saber la gama (mensual, trimestral, anual…) de cada OT ───\n'
             '   Generado por scratchpad/gen_gama_ciclos.py (IA17 + IP24 + IW49 + IP11).\n'
             '   GAMA_HR["RUTA/CONT"] = ciclos en días de los paquetes de la hoja de ruta (uno solo = gama segura).\n'
             '   GAMA_POS[posición] = [offsets posibles (meses), base meses, n de toma de referencia, fecha de esa toma, "RUTA/CONT"]:\n'
             '   la toma n lleva los paquetes cuyo ciclo (meses) divide a base·n + offset; si los offsets no coinciden, no se sabe. */\n')
    fh.write('const GAMA_HR = ' + json.dumps(hdr, separators=(',', ':')) + ';\n')
    fh.write('const GAMA_POS = ' + json.dumps(GAMA_POS, separators=(',', ':')) + ';\n')
print('hojas de ruta', len(hdr), '| un solo paquete', sum(1 for v in hdr.values() if len(v) == 1), '| posiciones con ciclo deducido', len(GAMA_POS))
