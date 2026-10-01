# -*- coding: utf-8 -*-
"""Genera assets/js/hidrolavado-sap-data.js: en qué mes programa SAP el hidrolavado de cada equipo.

El hidrolavado vive dentro de un paquete de la hoja de ruta (IA17). Cada plan con estrategia
(AASDF, día fijo) arma la toma n con los paquetes cuyo ciclo divide a base·n + offset (meses);
el offset depende de cómo se arrancó cada plan. Se deduce mirando qué paquete trajo de verdad
cada orden (texto de sus operaciones) y con eso se ubican las tomas que llevan el hidrolavado.

  HIDRO_SAP[equipo] = { hr, per, plan, meses:[m…] | cand:[[m…],[m…]], prox, ult }
    per   = cada cuántos meses SAP hace el hidrolavado (12 anual, 6 semestral)
    meses = meses del año (1-12) en que cae, si el historial alcanza para saberlo
    cand  = meses posibles cuando todavía no alcanza (planes nuevos / tomas sin operaciones)

Uso: python scratchpad/gen_hidrolavado_sap.py
Fuentes (Descargas):
  IP24 (1).xlsx                                   tomas por posición (fecha, orden, hoja de ruta)
  lista de operaciones.xlsx                       IW49: operaciones de órdenes abiertas
  Operaciones y órdenes de mantenimiento (1).xlsx operaciones (Fiori), órdenes ene-abr 2026
  Confirmaciones de orden de mantenimiento*.xlsx  operación notificada de órdenes ya cerradas
  + assets/js/hdr-plan-data.js (plan → equipo)
"""
import os, re, json, glob, collections, datetime as dt
import openpyxl

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
D = os.path.join(os.path.expanduser('~'), 'Downloads')
OUT = os.path.join(ROOT, 'assets', 'js', 'hidrolavado-sap-data.js')

# Hojas de ruta con hidrolavado: ciclos de sus paquetes (meses) y paquete que lo trae.
# conf=True: las órdenes traen solo la operación del paquete mayor (jerarquía), así que la operación
# notificada en Confirmaciones alcanza para saber el paquete (verificado: 315/315 iguales a IW49/Fiori).
# En los VRF la orden trae todas las operaciones y la confirmación figura siempre como "Bimestral"
# (56 de 68 no coinciden), así que ahí solo se usan IW49 y Fiori.
HOJAS = {
    'AACS5AEP/1': dict(ciclos=[1, 3, 12], hidro=12, conf=True, donde='MP Anual Roof Top R410A: "Hidrolavado serpentinas"'),
    'AACS4AEP/1': dict(ciclos=[1, 3, 12], hidro=12, conf=True, donde='MP Anual UTA: solo "tomar temperatura antes/luego de hidrolavar" (la tarea no está escrita)'),
    'AACS1AEP/1': dict(ciclos=[1, 3, 6], hidro=6, conf=True, donde='MP Semestral Roof-Top: "Hidrolavado de serpentinas de condensador y evaporador"'),
    'AAC19AEP/3': dict(ciclos=[2, 6, 12], hidro=6, conf=False, donde='MP Semestral VRF (paquetes 6M y 1A): "Hidrolavado de serpentinas de condensador"'),
    'AACS2AEP/4': dict(ciclos=[6], hidro=6, conf=True, donde='MP 6M Condensadora Multi Split: "Hidrolavado de condensador con aplicación de producto"'),
}
# Texto de la operación → ciclo (meses) del paquete que la trae
GAMA = [(r'anual|luego de hidrolavar', 12), (r'semestral', 6), (r'trimestral|verificar correcto funcionamiento de$', 3),
        (r'bimestral', 2), (r'mensual|^tomar temperatura de inyecci', 1)]


def gama_de(texto):
    t = texto.strip().lower()
    for rx, m in GAMA:
        if re.search(rx, t):
            return m
    return None


def num(s, rx=r'\((\d{9})\)\s*$'):
    m = re.search(rx, str(s or ''))
    return m.group(1) if m else None


# ── plan → equipo (vínculo exacto de Planes de Mantenimiento) ──
src = open(os.path.join(ROOT, 'assets', 'js', 'hdr-plan-data.js'), encoding='utf-8').read()
HDR_PLAN = json.loads(re.search(r'const HDR_PLAN = (\[.*?\]);\n', src, re.S).group(1))
equipo_de_plan = {r[1]: r[0] for r in HDR_PLAN}

# ── IP24: tomas de las posiciones cuya hoja de ruta trae hidrolavado ──
pos_info, tomas, orden_pos = {}, collections.defaultdict(list), {}
for r in openpyxl.load_workbook(os.path.join(D, 'IP24 (1).xlsx'), read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True):
    hr = f'{r[8]}/{r[9]}'
    if hr not in HOJAS or not isinstance(r[6], dt.datetime):
        continue
    pos = str(r[1])
    eq = equipo_de_plan.get(str(r[2]))
    if not eq:
        continue
    pos_info[pos] = dict(equipo=eq, hr=hr, plan=str(r[4] or ''))
    tomas[pos].append((int(r[5]), r[6].date(), str(r[7] or '')))
    if r[7]:
        orden_pos[str(r[7])] = pos

# ── qué operaciones trajo cada orden (3 fuentes) ──
ops = collections.defaultdict(set)
for r in openpyxl.load_workbook(os.path.join(D, 'lista de operaciones.xlsx'), read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True):
    if str(r[1]) in orden_pos:
        ops[str(r[1])].add(str(r[3] or ''))
for r in openpyxl.load_workbook(os.path.join(D, 'Operaciones y órdenes de mantenimiento (1).xlsx'), read_only=True, data_only=True).active.iter_rows(min_row=2, values_only=True):
    o = num(r[0])
    if o in orden_pos:
        ops[o].add(str(r[4] or ''))
for f in glob.glob(os.path.join(D, 'Confirmaciones de orden de mantenimiento*.xlsx')):
    it = openpyxl.load_workbook(f, read_only=True, data_only=True).active.iter_rows(values_only=True)
    h = next(it)
    io, iop = h.index('Orden'), h.index('Operación')
    for r in it:
        o = num(r[io])
        if o in orden_pos and HOJAS[pos_info[orden_pos[o]]['hr']]['conf']:
            ops[o].add(re.sub(r'\s*\(\d{4}\)\s*$', '', str(r[iop] or '')))


def due(m, ciclos):
    ds = [c for c in ciclos if m % c == 0]
    return max(ds) if ds else None


def sumar_meses(d, k):
    y, m = divmod(d.month - 1 + k, 12)
    return dt.date(d.year + y, m + 1, min(d.day, 28))


hoy = dt.date.today()
HIDRO_SAP, incons = {}, []
for pos, info in pos_info.items():
    cfg = HOJAS[info['hr']]
    cic, per = cfg['ciclos'], cfg['hidro']
    b = cic[0]
    T = sorted(tomas[pos])
    obs = []
    for n, f, o in T:
        gs = [g for g in (gama_de(t) for t in ops.get(o, ())) if g]
        if gs:
            obs.append((n, max(gs)))
    offs = [off for off in range(0, max(cic), b) if all(due(b * n + off, cic) == g for n, g in obs)]
    if not offs:
        incons.append(info['equipo'])
        continue
    # serie de tomas extendida 2 años después de la última (los planes nuevos traen pocas)
    last_n, last_f = T[-1][0], T[-1][1]
    serie = [(n, f) for n, f, _ in T] + [(last_n + k, sumar_meses(last_f, b * k)) for k in range(1, 24 // b + 1)]
    cands = []
    for off in offs:
        fechas = [f for n, f in serie if (b * n + off) % per == 0]
        cands.append((sorted({f.month for f in fechas}), fechas))
    uniq = {tuple(m) for m, _ in cands}
    rec = dict(hr=info['hr'], per=per, plan=info['plan'])
    if len(uniq) == 1:
        meses, fechas = cands[0]
        rec['meses'] = meses
        prox = [f for f in fechas if f >= hoy]
        ult = [f for f in fechas if f < hoy and f >= T[0][1]]
        if prox:
            rec['prox'] = prox[0].isoformat()
        if ult:
            rec['ult'] = ult[-1].isoformat()
    else:
        rec['cand'] = sorted(list(u) for u in uniq)
    prev = HIDRO_SAP.get(info['equipo'])
    if not prev or ('meses' in rec and 'meses' not in prev):
        HIDRO_SAP[info['equipo']] = rec

meta = dict(generado=hoy.isoformat(),
            fuentes=['IP24 (1).xlsx', 'lista de operaciones.xlsx', 'Operaciones y órdenes de mantenimiento (1).xlsx',
                     'Confirmaciones de orden de mantenimiento*.xlsx'],
            hojas={k: dict(per=v['hidro'], donde=v['donde']) for k, v in HOJAS.items()})
with open(OUT, 'w', encoding='utf-8') as fh:
    fh.write('/* ─── HIDRO_SAP — en qué mes programa SAP el hidrolavado de cada equipo ───\n'
             '   Generado por scratchpad/gen_hidrolavado_sap.py (IP24 + operaciones de cada orden: IW49,\n'
             '   Fiori y confirmaciones). HIDRO_SAP[equipo] = {hr, per (meses entre hidrolavados según SAP),\n'
             '   plan, meses [1-12] | cand [[…],[…]] si todavía no se puede saber, prox, ult}.\n'
             '   Lo usa el Calendario de Hidrolavados para anclar cada equipo al mes en que SAP ya\n'
             '   programa el MP que incluye el hidrolavado. */\n')
    fh.write('const HIDRO_SAP_META = ' + json.dumps(meta, ensure_ascii=False) + ';\n')
    fh.write('const HIDRO_SAP = ' + json.dumps(dict(sorted(HIDRO_SAP.items())), ensure_ascii=False, separators=(',', ':')) + ';\n')

seg = sum(1 for v in HIDRO_SAP.values() if 'meses' in v)
print(f'equipos {len(HIDRO_SAP)} | mes seguro {seg} | ambiguos {len(HIDRO_SAP) - seg} | inconsistentes {len(incons)} {incons[:10]}')
print(collections.Counter(v['hr'] for v in HIDRO_SAP.values()))
